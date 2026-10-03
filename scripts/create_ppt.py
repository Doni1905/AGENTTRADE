import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip('#')
    return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))

# Colors (White Theme)
BG_COLOR = hex_to_rgb("#ffffff")
FG_COLOR = hex_to_rgb("#172630")
ACCENT_COLOR = hex_to_rgb("#0e584f")
MUTED_COLOR = hex_to_rgb("#5b6c76")

prs = Presentation()

# Layouts
title_slide_layout = prs.slide_layouts[0]
blank_slide_layout = prs.slide_layouts[6]

def apply_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def add_title_slide(title_text, subtitle_text):
    slide = prs.slides.add_slide(title_slide_layout)
    apply_background(slide)
    
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = title_text
    title.text_frame.paragraphs[0].font.color.rgb = ACCENT_COLOR
    title.text_frame.paragraphs[0].font.name = 'Inter'
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.size = Pt(54)
    
    subtitle.text = subtitle_text
    subtitle.text_frame.paragraphs[0].font.color.rgb = MUTED_COLOR
    subtitle.text_frame.paragraphs[0].font.name = 'Inter'
    subtitle.text_frame.paragraphs[0].font.size = Pt(24)

def add_content_slide(title_text, bullet_points):
    slide = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide)
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.color.rgb = ACCENT_COLOR
    p.font.name = 'Inter'
    p.font.bold = True
    p.font.size = Pt(40)
    
    # Bullets
    body_box = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(5))
    tf = body_box.text_frame
    tf.word_wrap = True
    
    for pt in bullet_points:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.color.rgb = FG_COLOR
        p.font.name = 'Inter'
        p.font.size = Pt(28)
        p.space_before = Pt(14)

def add_architecture_diagram_slide():
    slide = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide)
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "System Architecture Diagram"
    p.font.color.rgb = ACCENT_COLOR
    p.font.name = 'Inter'
    p.font.bold = True
    p.font.size = Pt(40)
    
    def add_node(text, left, top, width, height, fill_color, border_color):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(2)
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.color.rgb = FG_COLOR
        p.font.name = 'Inter'
        p.font.size = Pt(18)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER
        return shape

    LIGHT_GREEN = hex_to_rgb("#e6f8f3")
    BORDER_GREEN = hex_to_rgb("#10b981")
    LIGHT_GRAY = hex_to_rgb("#f4f4f5")
    BORDER_GRAY = hex_to_rgb("#a1a1aa")
    
    u = add_node("User", 0.5, 3.5, 1.5, 1, LIGHT_GRAY, BORDER_GRAY)
    f = add_node("FastAPI Backend", 3, 3.5, 2, 1, LIGHT_GREEN, BORDER_GREEN)
    n = add_node("n8n Orchestrator", 6, 3.5, 2, 1, LIGHT_GREEN, BORDER_GREEN)
    a = add_node("Alpaca Trading", 3, 5.5, 2, 1, LIGHT_GRAY, BORDER_GRAY)
    q = add_node("Qdrant Vector DB", 6, 1.5, 2, 1, LIGHT_GREEN, BORDER_GREEN)
    o = add_node("Ollama (Qwen)", 8.5, 3.5, 1.5, 1, LIGHT_GREEN, BORDER_GREEN)
    
    # Simple lines connecting them
    def add_connector(x1, y1, x2, y2):
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        connector.line.color.rgb = BORDER_GRAY
        connector.line.width = Pt(2)

    add_connector(2, 4, 3, 4) # User to FastAPI
    add_connector(5, 4, 6, 4) # FastAPI to n8n
    add_connector(4, 4.5, 4, 5.5) # FastAPI to Alpaca
    add_connector(7, 3.5, 7, 2.5) # n8n to Qdrant
    add_connector(8, 4, 8.5, 4) # n8n to Ollama

add_title_slide("AGENTTRADE", "Multi-Agent Intelligence for Equity Research")

add_content_slide("Problem & Vision", [
    "Information overload in equity research",
    "Need for bounded, deterministic agent orchestration",
    "Live evidence + Local privacy + Agentic reasoning"
])

add_content_slide("System Architecture: Deep Dive", [
    "1. n8n Orchestrator: Bounded graph-based agent execution limits hallucinations and infinite loops.",
    "2. Qdrant Vector DB (RAG): Provides ultra-fast local retrieval of historical evidence cards and news.",
    "3. Local Ollama (Qwen 2.5): Executes complex reasoning (Bull/Bear cases, Critic judging) offline for privacy.",
    "4. FastAPI Backend: Manages state, routes traffic, and handles Alpaca paper trading integration.",
    "5. Frontend UI: Lightweight, responsive dashboard built with native web components and TradingView."
])

add_architecture_diagram_slide()

add_content_slide("Evaluation & Benchmarking", [
    "AD23731 Academic Standard Compliance",
    "Multi-tier model benchmarking (3B vs 7B)",
    "Strict bounding to prevent infinite agent loops"
])

add_content_slide("The Interface", [
    "Premium Dark Mode UI (Emerald & Obsidian)",
    "Live TradingView Advanced Charts",
    "Real-time research pipeline visibility"
])

prs.save('AgentTrade_Presentation.pptx')
print("Presentation saved as AgentTrade_Presentation.pptx")
